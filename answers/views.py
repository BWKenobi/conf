import os
import datetime

from datetime import date
from django.http import HttpResponse

from django.conf import settings
from django.core.files.storage import default_storage

from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib.auth import update_session_auth_hash

from django.template.loader import render_to_string
from django.core.mail import EmailMessage, send_mail

from django.contrib.auth.models import User

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.shared import Mm, Pt
from docx.enum.style import WD_STYLE_TYPE

from .models import Answer
from .forms import AnswerForm


@login_required(login_url='/login/')
def view_my_answers(request):
	username = request.user.username
	user = request.user

	my_answer = Answer.objects.filter(user = request.user).first()
	if not my_answer:
		my_answer = Answer.objects.create(user = request.user)

	form = AnswerForm(instance=my_answer, label_suffix='')

	if request.method=='POST':
		form = AnswerForm(request.POST, instance=my_answer, label_suffix='')

		if form.is_valid() and not my_answer.done:
			my_answer = form.save(False)
			my_answer.done = True
			my_answer.save()	

	args = {
		'my_answer': my_answer,
		'form': form
	}
	return render(request, 'answers/view_my_answers.html', args)


@login_required(login_url='/login/')
def get_answers(request):

	dte = date.today()
	document = Document()
	section = document.sections[-1]
	new_width, new_height = section.page_height, section.page_width
	section.orientation = WD_ORIENT.PORTRAIT
	section.page_width = Mm(210)
	section.page_height = Mm(297)
	section.left_margin = Mm(30)
	section.right_margin = Mm(10)
	section.top_margin = Mm(10)
	section.bottom_margin = Mm(10)
	section.header_distance = Mm(10)
	section.footer_distance = Mm(10)

	style = document.styles['Normal']
	font = style.font
	font.name = 'Times New Roman'
	font.size = Pt(12)

	answers = Answer.objects.filter(done = True).order_by('user__profile__surname')

	document.add_paragraph('Анкеты обратной связи').paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
	p = document.add_paragraph()
	p.add_run(dte.strftime('%d.%b.%Y')).italic = True
	p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.RIGHT

	for answer in answers:
		p = document.add_paragraph()
		p.add_run(answer.user.profile.get_full_name()).bold = True

		p = document.add_paragraph()
		p.add_run('Какая информация была наиболее интересной и полезной для  Вас?').italic = True
		p.paragraph_format.space_after = 0

		p = document.add_paragraph()
		p.add_run(answer.answer_1)

		p = document.add_paragraph()
		p.add_run('ККакие темы могут быть еще Вам интересны?').italic = True
		p.paragraph_format.space_after = 0

		p = document.add_paragraph()
		p.add_run(answer.answer_2)

		p = document.add_paragraph()
		p.add_run('Какие доклады и мастер-классы Вы могли бы предложить к проведению?').italic = True
		p.paragraph_format.space_after = 0

		p = document.add_paragraph()
		p.add_run(answer.answer_3)

		p = document.add_paragraph()

	response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
	response['Content-Disposition'] = 'attachment; filename=feedbacks (' + dte.strftime('%d-%b-%Y') + ').docx'
	document.save(response)

	return response


@login_required(login_url='/login/')
def get_answer(request):
	pk = request.GET['pk']

	if not request.user.profile.admin_access and not request.user.profile.moderator_access:
		return HttpResponse('', status = 500)

	answer = Answer.objects.filter(pk = pk, done = True).first()
	if not answer:
		return HttpResponse('', status = 500)

	return HttpResponse(render_to_string('answers/get_answer.html', {'answer': answer}))